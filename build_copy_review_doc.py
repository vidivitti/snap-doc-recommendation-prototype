import copy
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor


ROOT = Path('/Users/juliesutherland/Documents/ChatGPT/DTA SNAP PER + Efficiency Project (Julie)')
SOURCE_DOC = Path('/Users/juliesutherland/Documents/00-SNAP (DTA project) /Prototyping/Verification recommendation list pilot copy  - codex to generate and fill tables.docx')
SOURCE_APP = ROOT / 'src' / 'App.jsx'
OUTPUT_DOC = ROOT / 'Verification recommendation list pilot copy - populated for review.docx'


def extract_rules():
    source = SOURCE_APP.read_text(encoding='utf-8')
    start_token = 'const VERIFICATION_RULES = '
    end_token = ';\n\nconst SKIPPED_INCOME_RECOMMENDATIONS'
    start = source.index(start_token) + len(start_token)
    end = source.index(end_token, start)
    rules = json.loads(source[start:end])
    return [rule for rule in rules if not rule.get('suppressible', False)]


def clear_cell(cell):
    tc = cell._tc
    for paragraph in list(cell.paragraphs):
        tc.remove(paragraph._p)
    cell.add_paragraph()


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def add_hyperlink(paragraph, text, url, bold=False, italic=False):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True,
    )
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relationship_id)
    run = OxmlElement('w:r')
    run_properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '14558F')
    run_properties.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    run_properties.append(underline)
    if bold:
        run_properties.append(OxmlElement('w:b'))
    if italic:
        run_properties.append(OxmlElement('w:i'))
    run.append(run_properties)
    text_node = OxmlElement('w:t')
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_markdown_runs(paragraph, text, italic=False):
    pattern = re.compile(r'(\*\*.*?\*\*)')
    for part in pattern.split(text):
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        value = part[2:-2] if bold else part
        run = paragraph.add_run(value)
        run.bold = bold
        run.italic = italic


def add_segments(paragraph, segments, italic=False):
    for segment in segments:
        if isinstance(segment, str):
            run = paragraph.add_run(segment)
            run.italic = italic
            continue
        text = segment.get('strong') or segment.get('text') or ''
        if segment.get('href'):
            add_hyperlink(paragraph, text, segment['href'], bold='strong' in segment, italic=italic)
        else:
            run = paragraph.add_run(text)
            run.bold = 'strong' in segment
            run.italic = italic


def add_paragraph(cell, text='', *, bullet=False, bold=False, italic=False, space_after=4):
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = None
    paragraph.paragraph_format.keep_together = True
    if bullet:
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.add_run('• ')
    if bold:
        run = paragraph.add_run(text)
        run.bold = True
        run.italic = italic
    else:
        add_markdown_runs(paragraph, text, italic=italic)
    return paragraph


def write_plain_text(cell, text, italic=False):
    clear_cell(cell)
    first = True
    for block in str(text).split('\n\n'):
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        add_markdown_runs(paragraph, block.replace('\n', '\n'), italic=italic)
        paragraph.paragraph_format.space_after = None


def write_examples(cell, rule):
    clear_cell(cell)
    first = True
    for item in rule.get('examples', []):
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        paragraph.paragraph_format.space_after = None
        if isinstance(item, str):
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.add_run('• ')
            add_markdown_runs(paragraph, item)
            continue
        item_type = item.get('type')
        if item_type == 'heading':
            run = paragraph.add_run(item.get('text', ''))
            run.bold = True
        elif item_type == 'paragraph':
            add_segments(paragraph, item.get('segments', []))
        else:
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.add_run('• ')
            add_segments(paragraph, item.get('segments', []))
    note = rule.get('examplesNote')
    if note:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_before = None
        add_markdown_runs(paragraph, note)


def default_examples_label(rule):
    return f"Examples for {rule['title'].lower()}"


def fill_table(table, rule):
    examples_label = (
        'N/A — examples shown directly in the card body'
        if rule.get('examplesInline')
        else rule.get('examplesLabel') or default_examples_label(rule)
    )
    values = {
        0: rule['answer'],
        1: rule['title'],
        2: rule.get('microcopy', ''),
        3: examples_label,
        5: rule.get('details', {}).get('label', 'N/A'),
        6: rule.get('details', {}).get('content', 'N/A'),
        7: f"Rule ID: {rule['id']} | Source: {rule.get('source') or 'Not specified'}",
    }
    for row_index, row in enumerate(table.rows):
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        if row_index == 4:
            write_examples(row.cells[1], rule)
        else:
            write_plain_text(row.cells[1], values[row_index], italic=(row_index == 6 and rule.get('details', {}).get('italic', False)))
        for paragraph in row.cells[1].paragraphs:
            for run in paragraph.runs:
                run.font.highlight_color = None
                run.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
    if rule.get('learnMore'):
        cell = table.rows[2].cells[1]
        paragraph = cell.add_paragraph()
        add_hyperlink(paragraph, rule['learnMore']['text'], rule['learnMore']['href'])


def page_break_paragraph():
    paragraph = OxmlElement('w:p')
    run = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run.append(br)
    paragraph.append(run)
    return paragraph


def clone_and_insert(template_table, rules):
    fill_table(template_table, rules[0])
    last = template_table._tbl
    for rule in rules[1:]:
        page_break = page_break_paragraph()
        last.addnext(page_break)
        cloned = copy.deepcopy(template_table._tbl)
        page_break.addnext(cloned)
        table = template_table.__class__(cloned, template_table._parent)
        fill_table(table, rule)
        last = cloned


def main():
    source_rules = json.loads(
        SOURCE_APP.read_text(encoding='utf-8').split('const VERIFICATION_RULES = ', 1)[1].split(
            ';\n\nconst SKIPPED_INCOME_RECOMMENDATIONS', 1
        )[0]
    )
    rules = [rule for rule in source_rules if not rule.get('suppressible', False)]
    required = [rule for rule in rules if rule.get('required')]
    optional = [rule for rule in rules if not rule.get('required')]

    document = Document(SOURCE_DOC)
    if len(document.tables) != 3:
        raise RuntimeError(f'Expected 3 template tables; found {len(document.tables)}')

    required_template = document.tables[1]
    optional_template = document.tables[2]
    clone_and_insert(required_template, required)
    clone_and_insert(optional_template, optional)

    # Begin the recommendation review tables on a fresh page after the screen-copy section.
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == 'Required for decide if you still qualify':
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
            break

    # Keep the optional section label with its explanatory copy and begin it on a new page.
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == 'Optional to get a higher benefit amount':
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
            break

    document.core_properties.title = 'Recertification Verification Recommendation List Content Review'
    document.core_properties.subject = 'Draft recommendation copy for DTA stakeholder review'
    document.save(OUTPUT_DOC)
    print(json.dumps({
        'output': str(OUTPUT_DOC),
        'required_rules': len(required),
        'optional_rules': len(optional),
        'total_rules': len(rules),
        'excluded_suppressible_rules': len(source_rules) - len(rules),
    }, indent=2))


if __name__ == '__main__':
    main()
